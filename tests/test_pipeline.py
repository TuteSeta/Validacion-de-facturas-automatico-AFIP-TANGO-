import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import yaml
from openpyxl import Workbook, load_workbook

from src.main import run_validation
from src.transform import load_afip_result


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.arca = self.root / "arca.xlsx"
        self.odoo = self.root / "odoo.xlsx"
        self.output = self.root / "salida"
        self.config = self.root / "config.yaml"
        self._write_config()
        self._write_workbooks()

    def tearDown(self):
        self.temp.cleanup()

    def _write_config(self):
        config = {
            "origen_sheet": "ARCA",
            "destino_sheet": "IVA",
            "mapping": {
                "afip": {
                    "header_row": 2, "tipo": "Tipo", "pv": "PV", "num": "Numero",
                    "date": "Fecha", "cuit": "CUIT", "currency": "Moneda",
                    "exchange_rate": "TC",
                    "build_pattern": "{doc_type}-{letter}{pv:05d}{num:08d}",
                    "importes": {"exento": "Otros", "neto": "Neto", "iva": "IVA", "total": "Total"},
                },
                "odoo": {
                    "header_row": 3, "n_comp_column": "Comprobante", "date": "Fecha", "cuit": "CUIT",
                    "importes": {"exento": "Otros", "neto": "Neto", "iva": "IVA", "total": "Total"},
                },
            },
            "columns": [
                {"name": name, "type": "number", "tolerance": 0.01}
                for name in ("IMP_EXENTO", "IMP_NETO", "IMP_IVA", "IMP_TOTAL")
            ],
        }
        self.config.write_text(yaml.safe_dump(config, allow_unicode=True), encoding="utf-8")

    def _write_workbooks(self):
        arca = Workbook()
        ws = arca.active
        ws.title = "ARCA"
        ws.append(["Título preservado"])
        ws.append(["Fecha", "Tipo", "PV", "Numero", "CUIT", "Moneda", "TC", "Otros", "Neto", "IVA", "Total"])
        ws.append(["01/08/2026", "1 - Factura A", 1, 1, 30711111111, "$", 1, 0, 100, 21, 121])
        ws.append(["02/08/2026", "11 - Factura C", 1, 2, 30722222222, "$", 1, 0, 0, 0, 200])
        ws.append(["03/08/2026", "3 - Nota de Crédito A", 1, 3, 30733333333, "USD", 1000, 0, 10, 2.1, 12.1])
        ws.append(["03/08/2026", "2 - Nota de Débito A", 1, 4, 30744444444, "$", 1, 0, 100, 21, 121])
        ws.append(["03/08/2026", "1 - Factura A", 1, 5, "CUIT MALO", "$", 1, 0, 10, 2.1, 12.1])
        arca.create_sheet("Otra hoja")["A1"] = "conservar"
        arca.save(self.arca)

        odoo = Workbook()
        ws = odoo.active
        ws.title = "IVA"
        ws.append(["agosto"])
        ws.append([])
        ws.append(["Comprobante", "Fecha", "CUIT", "Otros", "Neto", "IVA", "Total"])
        ws.append(["FA-A 00001-00000001", "01/08/2026", 30711111111, 0, 100, 21, 121])
        ws.append(["FA-A 00001-00000009", "01/08/2026", 30799999999, 0, 50, 10.5, 60.5])
        ws.append(["FA-A 00001-00000010", "31/07/2026", 30799999998, 0, 10, 2.1, 12.1])
        ws.append(["NC-A 00001-00000003", "03/08/2026", 30733333333, 0, -10000, -2100, -12100])
        ws.append(["ND-A 00001-00000004", "03/08/2026", 30744444444, 0, 100, 21, 122])
        odoo.create_sheet("Filtros")["A1"] = "conservar"
        odoo.save(self.odoo)

    def test_pipeline_preserves_workbooks_and_reports_all_states(self):
        result = run_validation(
            str(self.arca), str(self.odoo), output_dir=str(self.output), config_path=str(self.config)
        )
        self.assertEqual(result["coincidencias"], 2)
        self.assertEqual(result["diferencias"], 1)
        self.assertEqual(result["faltantes_en_odoo"], 1)
        self.assertEqual(result["faltantes_en_arca"], 1)
        self.assertEqual(result["invalidos_arca"], 1)
        self.assertEqual(result["fuera_periodo_odoo"], 1)
        self.assertEqual(result["faltantes"], result["faltantes_en_odoo"])
        self.assertEqual(load_workbook(result["origen_validado"]).sheetnames, ["ARCA", "Otra hoja"])
        self.assertEqual(load_workbook(result["destino_validado"]).sheetnames, ["IVA", "Filtros"])
        report = load_workbook(result["reporte_validacion"], read_only=True)
        self.assertEqual(report.sheetnames, ["Resumen", "Conciliación", "Incidencias", "Odoo fuera de período"])
        self.assertTrue(Path(result["log"]).exists())

    def test_failure_does_not_publish_partial_run(self):
        with patch("src.main.write_consolidated_report", side_effect=OSError("fallo simulado")):
            with self.assertRaises(RuntimeError) as context:
                run_validation(
                    str(self.arca), str(self.odoo), output_dir=str(self.output), config_path=str(self.config)
                )
        self.assertIn("validacion_error_", str(context.exception))
        published = list(self.output.glob("Validacion_*"))
        self.assertEqual(published, [])
        self.assertEqual(len(list(self.output.glob("validacion_error_*.log"))), 1)

    def test_locked_output_is_reported_without_partial_run(self):
        with patch("src.main.write_odoo_validado", side_effect=PermissionError("archivo abierto")):
            with self.assertRaises(RuntimeError):
                run_validation(
                    str(self.arca), str(self.odoo), output_dir=str(self.output), config_path=str(self.config)
                )
        self.assertEqual(list(self.output.glob("Validacion_*")), [])

    def test_missing_required_column_creates_error_log(self):
        workbook = load_workbook(self.odoo)
        workbook["IVA"].cell(3, 7, "Total eliminado")
        workbook.save(self.odoo)
        with self.assertRaises(RuntimeError):
            run_validation(
                str(self.arca), str(self.odoo), output_dir=str(self.output), config_path=str(self.config)
            )
        self.assertEqual(len(list(self.output.glob("validacion_error_*.log"))), 1)

    def test_duplicate_arca_keys_are_excluded_and_reported(self):
        workbook = load_workbook(self.arca)
        values = [workbook["ARCA"].cell(3, column).value for column in range(1, 12)]
        workbook["ARCA"].append(values)
        workbook.save(self.arca)
        mapping = yaml.safe_load(self.config.read_text(encoding="utf-8"))["mapping"]
        loaded = load_afip_result(str(self.arca), "ARCA", mapping)
        duplicate_issues = loaded.issues[loaded.issues["MOTIVO"] == "clave de comprobante duplicada"]
        self.assertEqual(len(duplicate_issues), 2)


if __name__ == "__main__":
    unittest.main()
