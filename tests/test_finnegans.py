import tempfile
import unittest
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook

from src.main import run_validation
from src.transform import load_finnegans_result


class FinnegansTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.arca = self.root / "arca.xlsx"
        self.finnegans = self.root / "finnegans.xlsx"
        self.output = self.root / "salida"
        self.config = self.root / "config.yaml"
        self._write_config()
        self._write_workbooks()

    def tearDown(self):
        self.temp.cleanup()

    def _write_config(self):
        config = {
            "origen_sheet": "ARCA",
            "destino_sheet": "IVA",
            "finnegans_sheet": "hoja1",
            "mapping": {
                "afip": {
                    "header_row": 2,
                    "tipo": "Tipo",
                    "pv": "PV",
                    "num": "Numero",
                    "date": "Fecha",
                    "cuit": "CUIT",
                    "currency": "Moneda",
                    "exchange_rate": "TC",
                    "build_pattern": "{doc_type}-{letter}{pv:05d}{num:08d}",
                    "importes": {
                        "exento": "Otros", "neto": "Neto", "iva": "IVA", "total": "Total"
                    },
                },
                "finnegans": {
                    "header_row": 1,
                    "date": "Fecha",
                    "document": "Documento",
                    "n_comp_column": "Comprobante",
                    "cuit": "Cuit",
                    "importes": {
                        "exento": "Otros", "neto": "Imponible", "iva": "Impuesto", "total": "Total"
                    },
                },
            },
            "columns": [
                {"name": name, "type": "number", "tolerance": 0.01}
                for name in ("IMP_EXENTO", "IMP_NETO", "IMP_IVA", "IMP_TOTAL")
            ],
        }
        self.config.write_text(yaml.safe_dump(config, allow_unicode=True), encoding="utf-8")

    def _write_workbooks(self):
        arca = Workbook()
        ws = arca.active
        ws.title = "ARCA"
        ws.append(["Título"])
        ws.append(["Fecha", "Tipo", "PV", "Numero", "CUIT", "Moneda", "TC", "Otros", "Neto", "IVA", "Total"])
        ws.append(["01/09/2026", "1 - Factura A", 1, 1, 30711111111, "$", 1, 12, 100, 21, 133])
        ws.append(["30/09/2026", "3 - Nota de Crédito A", 2, 2, 30722222222, "$", 1, 0, 10, 2.1, 12.1])
        arca.save(self.arca)

        finnegans = Workbook()
        ws = finnegans.active
        ws.title = "hoja1"
        ws.append(["Fecha", "Documento", "Comprobante", "Proveedor", "Cuit", "Concepto", "Imponible", "Impuesto", "Otros", "Total"])
        ws.append(["01/09/2026", "FCCONTADO - 1", "A-00001-00000001", "Proveedor A", "30-71111111-1", "IVA CF", 100, 21, 0, 0])
        ws.append([None, None, None, None, None, "Exento", 0, 0, 12, 133])
        ws.append(["30/09/2026", "NCCPRA - 2", "A-00002-00000002", "Proveedor B", "30-72222222-2", "IVA CF", -10, -2.1, 0, -12.1])
        ws.append(["15/09/2026", "FC - 3", "C-00003-00000003", "Proveedor C", "30-73333333-3", "Exento", 0, 0, 50, 50])
        ws.append(["31/08/2026", "FC - 4", "A-00004-00000004", "Proveedor D", "30-74444444-4", "IVA CF", 10, 2.1, 0, 12.1])
        ws.append([None, None, "Total General", None, None, None, 100, 21, 62, 183])
        finnegans.save(self.finnegans)

    def test_loader_groups_detail_rows_and_ignores_totals(self):
        mapping = yaml.safe_load(self.config.read_text(encoding="utf-8"))["mapping"]
        loaded = load_finnegans_result(str(self.finnegans), "hoja1", mapping)

        self.assertEqual(len(loaded.records), 4)
        self.assertTrue(loaded.issues.empty)
        first = loaded.records.iloc[0]
        self.assertEqual(first["N_COMP"], "FA-A00001-00000001")
        self.assertEqual(first["FILAS_EXCEL"], (2, 3))
        self.assertEqual(first["IMP_EXENTO"], 12)
        self.assertEqual(first["IMP_NETO"], 100)
        self.assertEqual(first["IMP_IVA"], 21)
        self.assertEqual(first["IMP_TOTAL"], 133)
        self.assertEqual(loaded.records.iloc[1]["N_COMP"], "NC-A00002-00000002")
        self.assertEqual(loaded.records.iloc[1]["IMP_TOTAL"], -12.1)

    def test_finnegans_pipeline_and_multiline_marking(self):
        result = run_validation(
            str(self.arca),
            str(self.finnegans),
            output_dir=str(self.output),
            config_path=str(self.config),
            comparison_mode="finnegans",
        )

        self.assertEqual(result["modo"], "finnegans")
        self.assertEqual(result["coincidencias"], 2)
        self.assertEqual(result["diferencias"], 0)
        self.assertEqual(result["faltantes_en_finnegans"], 0)
        self.assertEqual(result["faltantes_en_arca"], 1)
        self.assertEqual(result["fuera_periodo_finnegans"], 1)
        self.assertEqual(result["finnegans_validado"], result["destino_validado"])

        workbook = load_workbook(result["finnegans_validado"])
        worksheet = workbook["hoja1"]
        self.assertEqual(worksheet.cell(2, 11).value, "Coincide")
        self.assertEqual(worksheet.cell(3, 11).value, "Coincide")
        report = load_workbook(result["reporte_validacion"], read_only=True)
        self.assertIn("Finnegans fuera de período", report.sheetnames)


if __name__ == "__main__":
    unittest.main()
