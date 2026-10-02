import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


def _excel_safe(frame):
    output = frame.copy()
    for column in output.columns:
        output[column] = output[column].map(
            lambda value: ", ".join(map(str, value)) if isinstance(value, (tuple, list)) else value
        )
    return output


def write_consolidated_report(path, reconciliation, source_names):
    metrics = reconciliation.metrics
    target_name = reconciliation.target_name
    target_key = target_name.casefold()
    summary = pd.DataFrame([
        ("Archivo ARCA", source_names["arca"]),
        (f"Archivo {target_name}", source_names[target_key]),
        ("Período desde", reconciliation.period_start.strftime("%d/%m/%Y")),
        ("Período hasta", reconciliation.period_end.strftime("%d/%m/%Y")),
        ("Coincidencias", metrics["coincidencias"]),
        ("Diferencias", metrics["diferencias"]),
        (f"Faltantes en {target_name}", metrics["faltantes_en_destino"]),
        ("Faltantes en ARCA", metrics["faltantes_en_arca"]),
        ("Inválidos ARCA", metrics["invalidos_arca"]),
        (f"Inválidos {target_name}", metrics["invalidos_destino"]),
        (f"Comprobantes {target_name} fuera del período", metrics["fuera_periodo_destino"]),
    ], columns=["Concepto", "Valor"])

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Resumen", index=False)
        _excel_safe(reconciliation.details).to_excel(
            writer, sheet_name="Conciliación", index=False
        )
        _excel_safe(reconciliation.issues).to_excel(
            writer, sheet_name="Incidencias", index=False
        )
        _excel_safe(reconciliation.out_of_period).to_excel(
            writer, sheet_name=f"{target_name} fuera de período", index=False
        )

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
        for column_cells in worksheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 60)
            worksheet.column_dimensions[column_cells[0].column_letter].width = width
    workbook.save(path)
